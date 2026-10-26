#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Assemble the Source Data workbook from the formal result tables.

Every sheet is built from formal results (including shared BEELINE case views), so that the Source Data file
shipped with the manuscript stays in step with the analysis outputs. The index
sheet describes each sheet by figure panel and content only; internal file
locations and intermediate artifacts are deliberately not recorded there.

Supplementary Tables S1 and S3 are typeset in the supplementary PDF and are
deliberately not repeated here.
"""

import os
import re
import sys
import subprocess
import tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).absolute().parents[1] / "functions"))
from runtime import use_python, task_environment, r_function_command
if __name__ == "__main__":
    use_python("general")

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = os.path.dirname(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
)
OUT_DIR = os.path.join(PROJECT_ROOT, "source_data")
SOURCE_XLSX = os.path.join(OUT_DIR, "Source Data.xlsx")
FONT_NAME = "Times New Roman"
HEADER_FILL = "D3D3D3"  # light gray
MAX_SHEET_NAME = 31

# Column headings keep these tokens in upper case.
_PRESERVE_UPPER = {
    "TF", "TFS", "CRE", "CRES", "RNA", "ATAC", "DNA", "GO", "BH", "FDR", "NES",
    "AP", "AUPRC", "AUROC", "EPR", "HSPC", "GSEA", "AUC", "AUCell", "UMAP",
    "PC", "PCA", "BIC", "RSS", "SNR", "ID", "IDS", "Jaccard", "Wilcoxon",
    "Holm", "KCNC2", "ZNF804A", "MK/E", "AUCELL", "SE",
}
_PRESERVE_MIXED = {
    "IN": "IN",
    "AUCell": "AUCell", "AUPRC": "AUPRC", "AUROC": "AUROC", "NES": "NES",
    "EPR": "EPR", "HSPC": "HSPC", "GSEA": "GSEA", "Jaccard": "Jaccard",
    "Wilcoxon": "Wilcoxon", "Holm": "Holm", "Benjamini": "Benjamini",
    "Hochberg": "Hochberg", "TF": "TF", "TFs": "TFs", "CRE": "CRE",
    "CRES": "CREs", "RNA": "RNA", "ATAC": "ATAC", "DNA": "DNA", "GO": "GO",
    "BH": "BH", "FDR": "FDR", "AP": "AP", "AUC": "AUC", "BIC": "BIC",
    "UMAP": "UMAP",
}


def log(message, level="info"):
    prefix = {"info": "[info]", "warn": "[warn]", "fail": "[fail]", "success": "[ok]"}.get(level, f"[{level}]")
    print(f"{prefix} {message}", flush=True)


def _format_column_name(name):
    """Turn a result-table column name into a readable heading."""
    key = str(name).strip()
    if key in _PRESERVE_MIXED:
        return _PRESERVE_MIXED[key]
    words = re.split(r"[_\s]+", key)
    out = []
    for i, w in enumerate(words):
        if not w:
            continue
        if w in _PRESERVE_MIXED:
            out.append(_PRESERVE_MIXED[w])
        elif w.upper() in _PRESERVE_UPPER:
            out.append(w.upper())
        elif w.lower() in ("of", "and", "per", "vs", "by", "in", "to") and i:
            out.append(w.lower())
        else:
            out.append(w[:1].upper() + w[1:])
    return " ".join(out) if out else key


def _format_columns(df):
    df = df.copy()
    df.columns = [_format_column_name(c) for c in df.columns]
    return df


def _sanitize_sheet_name(name):
    clean = re.sub(r"[\[\]:*?/\\]", "-", name).strip()
    if len(clean) > MAX_SHEET_NAME:
        clean = clean[:MAX_SHEET_NAME]
    return clean


def read_table(relative_path):
    """Read a result table (tsv/csv, optionally gzip-compressed)."""
    path = os.path.join(PROJECT_ROOT, relative_path)
    if path.endswith(".gz"):
        engine_path = path
        lowered = path[:-3].lower()
    else:
        engine_path = path
        lowered = path.lower()
    sep = "," if lowered.endswith(".csv") else "\t"
    return pd.read_csv(engine_path, sep=sep, compression="infer", low_memory=False)


# --------------------------------------------------------------------------
# Source Data sheets: one entry per formal result table.
#   figure  : figure panel(s) the table supports
#   content : what the table contains, in article terms
# --------------------------------------------------------------------------
SOURCE_SHEETS = [
    # Figure 2 - case-matched BEELINE evaluation
    dict(sheet="BEELINE_case_metrics", figure="Figure 2a,b; Figures S1-S3",
         content="Complete case-level AP, benchmark metrics, prevalence, within-case ranks and paired AP differences (multiCSN minus comparator) for 500 cases and 11 methods; self-comparisons are blank; DatasetType distinguishes curated and synthetic cases",
         path="results/beeline/metrics/case_metrics.tsv", transform="beeline_metrics"),
    dict(sheet="BEELINE_AP_summary", nonempty=True, figure="Figure S1; Table S2",
         content="Mean AP and trapezoidal AUPRC, with case counts and mean case ranks for all 500 cases and the curated and synthetic subsets",
         path="results/beeline/ap_sensitivity/method_summary.tsv"),
    dict(sheet="BEELINE_AP_groups", figure="Figure S1",
         content="AP and trapezoidal AUPRC means by benchmark model and method",
         path="results/beeline/ap_sensitivity/group_summary.tsv"),
    dict(sheet="BEELINE_AP_paired_summary", nonempty=True, figure="Figure S1",
         content="Descriptive paired AP differences and case win/tie/loss counts",
         path="results/beeline/ap_sensitivity/paired_summary.tsv"),
    dict(sheet="Fig2b_rank_summary", figure="Figure 2b",
         content="Mean, median and quartile within-case ranks per method and metric across the 500 cases",
         path="results/beeline/metrics/rank_summary.tsv"),
    dict(sheet="Fig2c_runtime_summary", figure="Figure 2c",
         content="Single-worker inference-and-export runtime totals, medians, maxima and ratios to multiCSN (500 cases per method; input reading and matrix preparation excluded)",
         path="results/beeline/runtime/runtime_summary.tsv"),
    dict(sheet="Fig2c_runtime_all_methods", figure="Figure 2c",
         content="Per-case inference-and-export runtimes for 11 methods; input reading and matrix preparation are excluded",
         path="results/beeline/runtime/runtime_all_methods.tsv"),
    dict(sheet="Fig2d_HSC_truth", figure="Figure 2d",
         content="Reference network for the HSC 2000-3-30 BEELINE case",
         case="HSC-2000-3-30", view="truth"),
    dict(sheet="Fig2d_HSC_methods", figure="Figure 2d",
         content="Per-method edge counts, precision, recall, F1 and sign concordance for the HSC 2000-3-30 case",
         case="HSC-2000-3-30", view="summary"),
    dict(sheet="Fig2d_HSC_edges", figure="Figure 2d",
         content="Displayed edges of every method for the HSC 2000-3-30 case at the specified edge budget",
         case="HSC-2000-3-30", view="edges"),

    # Figure 3 - hematopoietic Perturb-multiome
    dict(sheet="Fig3a_response_metrics", figure="Figure 3a",
         content="Equal-TF mean paired differences, TF-bootstrap intervals and adjusted P values",
         path="source_data/files/fig3a_multimetric_effects_source.tsv"),
    dict(sheet="Fig3a_TF_paired_effects", figure="Figure 3a",
         content="TF-level mean paired differences and finite timepoint counts underlying equal-TF effects and intervals",
         path="source_data/files/fig3a_tf_cluster_effects_source.tsv"),
    dict(sheet="Fig3b_TF_region", figure="Figure 3b",
         content="UniBind TF-region precision and recall in the four-method common non-empty strata",
         path="source_data/files/fig3b_tf_region_source.tsv"),
    dict(sheet="Fig3c_region_gene", figure="Figure 3c",
         content="Target-blocked region-gene contact enrichment over distance-matched expectation",
         path="results/gse274113/comparison/region_gene_contact_four_methods.tsv"),
    dict(sheet="Fig3d_direct_enrichment", figure="Figure 3d",
         content="Direct RNA- and ATAC-target response enrichment per method, regulator and time point",
         path="source_data/files/fig3d_direct_enrichment_source.tsv.gz"),
    dict(sheet="Fig3ef_RUNX1_paths", figure="Figure 3e,f",
         content="RUNX1 day-9 direction-consistent direct and second-order paths",
         path="results/gse274113/neighborhood/hero_direction_consistent_paths.tsv"),
    dict(sheet="Fig3ef_chain_annotations", figure="Figure 3e,f,j",
         content="All selected chains across TFs, methods and four time points, with RNA/ATAC effects, FDR and direction support; Fig. 3f uses jointly evaluable RUNX1 day-9 rows",
         path="results/gse274113/comparison/selected_chain_response_annotations.tsv.gz",
         columns=["method", "timepoint", "perturbation_regulator", "regulator", "region", "target",
                  "chain_direction", "tf_region_direction", "region_gene_direction",
                  "atac_truth_feature", "atac_logFC", "atac_FDR", "rna_logFC", "rna_FDR",
                  "joint_evaluable", "gene_direction_correct", "atac_direction_correct",
                  "mediated_direction_correct"]),
    dict(sheet="Fig3g_response_order", figure="Figure 3g",
         content="Direct and second-order RNA-response enrichment per perturbed regulator",
         path="source_data/files/fig3g_order_response_source.tsv"),
    dict(sheet="Fig3h_concordance", figure="Figure 3h",
         content="Signed response concordance per method, regulator and time point",
         path="source_data/files/fig3h_response_concordance_source.tsv"),
    dict(sheet="Fig3h_summary", figure="Figure 3h",
         content="Median and interquartile signed response concordance per method and time point",
         path="source_data/files/fig3h_response_concordance_summary.tsv"),
    dict(sheet="Fig3i_layer_attrition", figure="Figure 3i",
         content="Replicate-deletion retention of complete-network edges by eligibility layer and fold",
         path="results/gse274113/stability/layer_attrition.tsv"),
    dict(sheet="Fig3j_chain_response", figure="Figure 3j",
         content="Joint ATAC-RNA response support of selected chains per method and TF cluster",
         path="results/gse274113/comparison/chain_response_four_methods.tsv"),

    # Figure 4 - developing human fetal cortex
    dict(sheet="Fig4a_topology", figure="Figure 4a",
         content="Candidate-pair and selected-edge counts for the three component networks",
         path="results/gse162170/network/network_topology.tsv"),
    dict(sheet="Fig4b_target_funnel", figure="Figure 4b",
         content="Target funnel from annotated genes to selected sparse support",
         path="results/gse162170/network/Fig4b_target_funnel.tsv"),
    dict(sheet="Fig4c_chain_summary", figure="Figure 4c",
         content="Exact-region chain counts and direct TF-gene overlap",
         path="results/gse162170/network/chain_summary.tsv"),
    dict(sheet="Fig4d_terminal_states", figure="Figure 4d",
         content="Mutually exclusive terminal target classes",
         path="results/gse162170/network/terminal_states.tsv"),
    dict(sheet="Fig4e_aucell_by_class", figure="Figure 4e",
         content="RNA AUCell of the five chromatin-anchored target programs by cell class",
         path="results/gse162170/cell_programs/aucell_by_class.tsv"),
    dict(sheet="Fig4e_aucell_by_cell", nonempty=True, figure="Figure 4e",
         content="All five RNA AUCell scores per cell, UMAP coordinates and cell annotations",
         path="results/gse162170/cell_programs/aucell_by_cell.tsv.gz",
         transform="aucell_cells"),
    dict(sheet="Fig4e_threshold_summary", nonempty=True, figure="Figure 4e",
         content="Fixed program-specific AUC thresholds, above-threshold counts, cell totals and fractions by cell class",
         path="results/gse162170/cell_programs/aucell_threshold_summary.tsv"),
    dict(sheet="Fig4f_module_activity", figure="Figure 4f",
         content="Aligned RNA target-AUCell and selected-region ATAC activity for the 25 regulator-anchor-sign modules",
         path="results/gse162170/cell_programs/module_raw_and_scaled_activity.tsv"),
    dict(sheet="Fig4f_S9_module_summary", figure="Figure 4f; Figure S9",
         content="Membership and size summary of the 25 regulator-anchor-sign modules",
         path="results/gse162170/cell_programs/module_summary.tsv"),
    dict(sheet="Fig4g_GO_all_tests", figure="Figure 4g",
         content="All tested GO BP terms, raw P values and within-family/global BH adjustments; Displayed marks the plotted terms; 11017 input targets and 7613 GO-annotated background genes",
         path="results/gse162170/cell_programs/go_all_terms.tsv.gz", transform="cortex_go_display"),
    dict(sheet="Fig4h_gsea_terms", figure="Figure 4h",
         content="Displayed preranked GSEA terms over the selected-target RNA rank",
         path="results/gse162170/key_gene_gsea/gsea_selected_terms.tsv"),
    dict(sheet="Fig4h_gsea_genes", figure="Figure 4h",
         content="Leading-edge and rank statistics of the displayed GSEA terms",
         path="results/gse162170/key_gene_gsea/gsea_key_genes.tsv"),
    dict(sheet="Fig4h_rank_curve", nonempty=True, figure="Figure 4h",
         content="Complete 11,017-target IPC/RG RNA ranking, exact plotted running enrichment score, pathway membership and leading-edge membership",
         path="results/gse162170/key_gene_gsea/gsea_rank_curve.tsv.gz"),
    dict(sheet="Fig4ij_S11_ZNF804A_paths", figure="Figure 4i,j; Figure S11a,b",
         content="The 32 ZNF804A exact-region paths with regulators, regions, signs and deletion evidence",
         path="results/gse162170/loci/candidates/ZNF804A_paths.tsv",
         identity=("target", "ZNF804A")),
    dict(sheet="Fig4i_ZNF804A_annotation", figure="Figure 4i",
         content="ZNF804A gene annotation used with the path coordinates to draw the locus track",
         path="results/gse162170/loci/candidates/ZNF804A_gene_annotation.tsv",
         identity=("gene_name", "ZNF804A")),
    dict(sheet="FigS11a_ZNF804A_expression", figure="Figure S11a",
         content="Per-sample ZNF804A RNA expression by cell class",
         path="results/gse162170/loci/candidates/ZNF804A_expression_samples.tsv",
         identity=("target", "ZNF804A")),
    dict(sheet="Fig4i_ZNF804A_accessibility", figure="Figure 4i",
         content="Region accessibility by cell class for the ZNF804A locus track",
         path="results/gse162170/loci/candidates/ZNF804A_accessibility.tsv"),
    dict(sheet="FigS11_KCNC2_paths", figure="Figure S11c-e",
         content="The 33 KCNC2 exact-region paths with regulators, regions, signs and deletion evidence",
         path="results/gse162170/loci/candidates/KCNC2_paths.tsv",
         identity=("target", "KCNC2")),
    dict(sheet="FigS11c_KCNC2_annotation", figure="Figure S11c",
         content="KCNC2 gene annotation used with the path coordinates to draw the locus track",
         path="results/gse162170/loci/candidates/KCNC2_gene_annotation.tsv",
         identity=("gene_name", "KCNC2")),
    dict(sheet="FigS11d_KCNC2_expression", figure="Figure S11d",
         content="Per-sample KCNC2 RNA expression by cell class",
         path="results/gse162170/loci/candidates/KCNC2_expression_samples.tsv",
         identity=("target", "KCNC2")),
    dict(sheet="FigS11c_KCNC2_accessibility", figure="Figure S11c",
         content="Region accessibility by cell class for the KCNC2 locus track",
         path="results/gse162170/loci/candidates/KCNC2_accessibility.tsv"),

    # Figure 5 - human erythroid maturation
    dict(sheet="Fig5a_cells_pseudotime", figure="Figure 5a",
         content="Current plotted RNA UMAP and expression pseudotime of all 10,952 erythroid cells",
         path="results/gse194122/trajectory/erythroid_gex_umap.tsv.gz", transform="erythroid_coordinates"),
    dict(sheet="Fig5b_g_focal_regulons", figure="Figure 5b,g",
         content="The 6 focal regulons with selection statistics and HSC-to-Normoblast activity change",
         path="results/gse194122/programs/hero_regulons.tsv"),
    dict(sheet="Fig5b_regulon_network", figure="Figure 5b",
         content="Global map of the six focal positive regulons with donor-stable chain support",
         path="results/gse194122/programs/hero_regulon_network.tsv.gz"),
    dict(sheet="Fig5c_state_activity", figure="Figure 5c",
         content="State-aggregated gene-module expression and CRE-linked network support per regulon and state",
         path="results/gse194122/programs/module_state_summary.tsv.gz"),
    dict(sheet="Fig5d_feature_programs", figure="Figure 5d",
         content="Temporal cluster assignment of the 446 focal TF and target genes",
         path="source_data/files/dynamic_feature_programs.tsv"),
    dict(sheet="Fig5d_go_enrichment", figure="Figure 5d",
         content="GO biological-process enrichment per temporal cluster of the fitted trajectories",
         path="source_data/files/dynamic_go_enrichment.tsv.gz"),
    dict(sheet="Fig5e_h_model_variance", figure="Figure 5e,h",
         content="Per-regulon activity coupling and partial state and donor R-squared values",
         path="results/gse194122/programs/module_model_variance.tsv.gz"),
    dict(sheet="Fig5f_transition_topology", figure="Figure 5f",
         content="Gained, lost, sign-stable and sign-flipped support across adjacent erythroid states",
         path="results/gse194122/trajectory/GSE5d_adjacent_transition_topology.tsv"),

    dict(sheet="Fig5e_budget_coupling", nonempty=True, figure="Figure 5e",
         content="Ten equal-budget donor-exclusion reconstructions: finite module counts and median five-state RNA-support correlations",
         path="results/gse194122/cell_budget/global_coupling_summary.tsv"),
    dict(sheet="Erythroid_target_retention", nonempty=True, figure="Methods: Erythroid network stability and program analysis",
         content="Complete HMGA1 and BNC2 direct-target retention in ten 500-cell donor-exclusion reconstructions",
         path="results/gse194122/cell_budget/focal_target_set_retention.tsv"),
    dict(sheet="Erythroid_GO_sensitivity", nonempty=True, figure="Methods: Erythroid network stability and program analysis",
         content="Retention of all original HMGA1 and BNC2 GO terms with same-family full-cohort controls",
         path="results/gse194122/cell_budget/displayed_GO_reproducibility.tsv"),
    dict(sheet="TableS6_GATA1_gate", nonempty=True, figure="Table S6",
         content="Five-state GATA1 cross-cohort assessment combining reconstruction consensus, ChIP overlap and RNA responses",
         path="results/gse194122/cases/GATA1_gate_stages.tsv"),

    # Supplementary Figures
    dict(sheet="FigS4_scaling_runs", figure="Figure S4",
         content="Runtime and sampled process-tree RSS sum for 60 fits across cell counts, total gene counts and specified cores",
         path="results/beeline/core_scaling/core_scaling_runs.csv",
         columns=["Case", "Scenario", "AxisValue", "AxisLabel", "Cores", "InputCells", "InputGenes",
                  "RuntimeSeconds", "UserSeconds", "SystemSeconds", "OutputEdges",
                  "MemoryTreePeakMb", "MemoryTreeBeforeMb", "MemoryTreeProcessesAtPeak", "MemorySamples",
                  "MemorySamplingIntervalSeconds", "MemoryMaxObservedIntervalSeconds",
                  "TimingScope"]),
    dict(sheet="FigS4_core_comparison", figure="Figure S4",
         content="One-core versus ten-core inference runtime and peak RSS sum (MiB) per scaling configuration with selected-edge counts",
         path="results/beeline/core_scaling/one_to_ten_core_comparison.csv"),
    dict(sheet="FigS4_scaling_summary", figure="Figure S4",
         content="Summary speedup and memory ratios of the scaling experiment",
         path="results/beeline/core_scaling/scaling_summary_statistics.csv"),
    dict(sheet="FigS4_edge_invariance", figure="Figure S4",
         content="Selected-edge counts across specified core settings",
         path="results/beeline/core_scaling/edge_count_invariance.csv",
         columns=["Case", "CoreSettings", "EdgeCounts", "OutputEdges"]),
    dict(sheet="FigS5_VSC_truth", figure="Figure S5",
         content="Reference network for the VSC 2000-9 BEELINE case",
         case="VSC-2000-9", view="truth"),
    dict(sheet="FigS5_VSC_methods", figure="Figure S5",
         content="Per-method recovery and sign concordance for the VSC 2000-9 case",
         case="VSC-2000-9", view="summary"),
    dict(sheet="FigS5_VSC_edges", figure="Figure S5",
         content="Displayed edges of every method for the VSC 2000-9 case",
         case="VSC-2000-9", view="edges"),
    dict(sheet="FigS6_scMultiSim_seeds", figure="Figure S6a-c",
         content="Complete per-seed TF-gene and region-gene recovery metrics, ranking diagnostics and prevalence baselines across the 10 scMultiSim simulations",
         path="results/scmultisim/multiseed/per_seed_diagnostics.tsv"),
    dict(sheet="FigS6_truth_sensitivity", figure="Figure S6c",
         content="Median paired static-input-GRN minus dynamic-GRN-union recovery differences across seeds",
         path="results/scmultisim/multiseed/static_dynamic_truth_sensitivity.tsv"),
    dict(sheet="FigS6_ablation_seeds", figure="Figure S6d-f",
         content="Per-seed TF-gene recovery for the forward-only and full add-delete-swap search variants",
         path="results/scmultisim/component_ablation/endpoint_metrics_all_seeds.tsv"),
    dict(sheet="FigS7_runtime", figure="Figure S7a (runtime)",
         content="Measured runtime of the four multi-omic methods on GSE274113, with the core count and method-specific timing scope of each configuration",
         path="results/gse274113/runtime/multiomic_runtime_comparison.tsv"),
    dict(sheet="FigS7_runtime_timepoints", figure="Figure S7a (runtime)",
         content="Four-timepoint timings for the one-core and eight-core configurations, with overall and stage-specific runtimes",
         path="results/gse274113/runtime/multiomic_runtime_timepoints.tsv"),
    dict(sheet="FigS7_fold_stability", figure="Figure S7c",
         content="Per-fold replicate-deletion stability metrics",
         path="results/gse274113/stability/fold_decomposition.tsv"),
    dict(sheet="FigS7_timepoint_summary", figure="Figure S7 (stability summary)",
         content="Per-time-point medians and quartiles of replicate-deletion stability metrics",
         path="results/gse274113/stability/timepoint_stability_summary.tsv"),
    dict(sheet="FigS7_pairwise_stability", figure="Figure S7c",
         content="Edge-set Jaccard similarity for fold pairs and fold-versus-complete comparisons",
         path="results/gse274113/stability/pairwise_fold_metrics.tsv"),
    dict(sheet="FigS7_rank_deciles", figure="Figure S7e",
         content="Edge retention by within-network rank decile",
         path="results/gse274113/stability/rank_decile_decomposition.tsv"),
    dict(sheet="FigS7_top100", figure="Figure S7d",
         content="Top-100 edge rank decomposition of replicate-deletion retention",
         path="results/gse274113/stability/top100_decomposition.tsv"),
    dict(sheet="FigS7b_endpoint_counts", figure="Figure S7b",
         content="Selected TF-gene, TF-region, region-gene and exact-chain counts at each time point",
         path="results/gse274113/chains/timepoint_summary.tsv"),
    dict(sheet="FigS8_strata", figure="Figure S8",
         content="Per-stratum direct and chain-representation response metrics",
         path="results/gse274113/ablation/stratum_metrics.tsv"),
    dict(sheet="FigS8_matched_budgets", figure="Figure S8c",
         content="Matched-budget precision, common-support eligibility and TF/timepoint identities",
         path="results/gse274113/ablation/matched_budget_metrics.tsv"),
    dict(sheet="FigS8_TF_metrics", figure="Figure S8a,b",
         content="TF-level mean AP and response coverage for each representation",
         path="results/gse274113/ablation/TF_macro_metrics.tsv"),
    dict(sheet="FigS8_budget_summary", figure="Figure S8c",
         content="Matched-budget summary statistics",
         path="results/gse274113/ablation/matched_budget_summary.tsv"),
    dict(sheet="FigS8_score_resolution", figure="Figure S8",
         content="Score resolution of direct and chain representations",
         path="results/gse274113/ablation/score_resolution.tsv"),
    dict(sheet="FigS9_module_membership", figure="Figure S9",
         content="Target membership of the 25 regulator-anchor-sign modules",
         path="results/gse162170/cell_programs/module_membership.tsv.gz"),
    dict(sheet="FigS12_network_overlap", figure="Figure S12a",
         content="Within-cell-class versus pooled direct-edge and exact-region-chain overlap",
         path="results/gse162170/within_class/pooled_within_class_overlap.tsv"),
    dict(sheet="FigS12_locus_recovery", figure="Figure S12b",
         content="Same-sign recovery fractions and chain-count labels for both focal loci",
         path="results/gse162170/within_class/focal_locus_overlap.tsv"),
    dict(sheet="FigS12_chain_retention", figure="Figure S12b",
         content="Individual pooled focal chains and their recovery within each cell class",
         path="results/gse162170/within_class/focal_chain_retention.tsv"),
    dict(sheet="FigS14_regulator_turnover", figure="Figure S14",
         content="Tie-inclusive TF-gene and factorized-chain regulator turnover across adjacent erythroid states",
         path="results/gse194122/trajectory/GSE5e_regulator_transition_counts_display.tsv"),
    dict(sheet="FigS13_cell_budget_a", figure="Figure S13a",
         content="Signed Jaccard similarity between each fixed-cell-budget fit and the all-cell fit",
         path="source_data/files/figs13_panel_a_state_overlap.tsv",
         display_labels={"endpoint_label": {"TF-gene": "TF–gene"}}),
    dict(sheet="FigS13_cell_budget_b", figure="Figure S13b",
         content="Adjacent-state topology-composition similarity under fixed-cell-budget refitting",
         path="source_data/files/figs13_panel_b_transition_similarity.tsv",
         display_labels={"endpoint_label": {"TF-gene": "TF–gene"},
                         "transition_label": {"HSC to MK/E": "HSC → MK/E", "MK/E to ProE": "MK/E → ProE",
                                              "ProE to Ery": "ProE → Ery", "Ery to Norm": "Ery → Norm"}}),
    dict(sheet="Erythroid_target_support", figure="Supplementary Table S5; Source Data",
         content="Target-deduplicated focal-module support retained without a separate plot",
         path="source_data/files/erythroid_focal_target_support.tsv"),
    dict(sheet="FigS13_cell_budget_c", figure="Figure S13c",
         content="Selected exact-region chain counts divided by fixed consensus-module target counts",
         path="source_data/files/figs13_panel_c_focal_chain_support.tsv"),
    dict(sheet="BEELINE_case_manifest", figure="Figure 2; Figures S1-S3",
         content="Model, cells, dropout, replicate and gene/truth-edge counts from the reference networks",
         path="results/beeline/evaluation/beeline_case_manifest.tsv"),
    dict(sheet="TableS5_budget_summary", figure="Table S5",
         content="Focal-module peak and expression-support correlation sensitivity; RNA summaries held fixed",
         path="results/gse194122/cell_budget/focal_budget_summary.tsv"),
    dict(sheet="TableS5_budget_draws", figure="Table S5",
         content="Per-draw peak sets and correlations for target and exact-chain support",
         path="results/gse194122/cell_budget/focal_budget_draws.tsv"),
    dict(sheet="FigS10a_b_locus_screen", figure="Figure S10a,b",
         content="Complete ordered fetal-cortex locus screen with all screening statistics and within-class ranks",
         path="results/gse162170/loci/eligible_loci_by_class.tsv"),

    dict(sheet="Fig5ij_RNA_donor_scores", nonempty=True, figure="Figure 5i,j; Figure S15c,d",
         content="All four fixed functional subsets and TF RNA, feature-standardized over 49 donor-state observations; no RNA-based gene filtering",
         path="results/gse194122/cases/program_donor_RNA_scores.tsv"),
    dict(sheet="Fig5ij_S15_program_members", nonempty=True, figure="Figure 5i,j; Figure S15b-d",
         content="Complete GO-defined display memberships for HMGA1 (16), SOX6 (14), KLF1 (13) and BNC2 (39)",
         path="results/gse194122/cases/program_members.tsv"),
    dict(sheet="FigS15a_HMGA1_retention", nonempty=True, figure="Figure S15a",
         content="Reference recall and budget containment for all ten HMGA1 loss-set reconstructions; full-cohort reference also retained",
         path="results/gse194122/cases/HMGA1_target_retention.tsv"),
    dict(sheet="FigS15b_SOX6_link_recurrence", nonempty=True, figure="Figure S15b",
         content="All 14 SOX6 functional-subset genes across five states; same-sign direct-link selection counts out of ten and reference selection",
         path="results/gse194122/cases/SOX6_signed_link_recurrence.tsv"),
    dict(sheet="Erythroid_reported_GO", nonempty=True, figure="Figure 5i,j; Figure S15c,d",
         content="Reported GO tests with explicit endpoint and correction family; HMGA1 GO test is for the 23-gene core, whose 16 translation members are plotted",
         path="results/gse194122/cases/reported_program_GO.tsv"),
    dict(sheet="Fig5ij_GO_display", nonempty=True, figure="Figure 5i,j",
         content="Five displayed GO terms with full tested set sizes, term overlaps, explicit backgrounds and distinct BH correction families",
         path="source_data/files/fig5ij_GO_display.tsv"),
]



_case_views = {}


def transform_source(df, name):
    if name == "beeline_metrics":
        keys = ["Case", "Method"]
        extra = read_table("results/beeline/metrics/combined__other_metrics.csv")
        metrics = ["AUROC", "ACC", "Precision", "Recall", "F1"]
        df = df.drop(columns=["Dataset", "ReferenceNetwork", "Species"])
        df = df.merge(extra[keys + metrics], on=keys, how="left", validate="one_to_one", sort=False)
        ranks = read_table("results/beeline/metrics/case_ranks.tsv")
        ranks = ranks.pivot(index=keys, columns="Metric", values="Rank").add_suffix("_rank")
        df = df.merge(ranks, on=keys, how="left", validate="one_to_one", sort=False)
        ap = read_table("results/beeline/ap_sensitivity/case_metrics.tsv")
        additional = [column for column in ap if column not in df.columns]
        df = df.merge(ap[keys + additional], on=keys, how="left", validate="one_to_one", sort=False)
        paired = read_table("results/beeline/ap_sensitivity/paired_differences.tsv")
        paired = paired[keys + ["Difference"]].rename(columns={"Difference": "Paired_AP_difference"})
        return df.merge(paired, on=keys, how="left", validate="one_to_one", sort=False)
    if name == "cortex_go_display":
        selected = read_table("results/gse162170/cell_programs/go_display_terms.tsv")
        keys = ["anchor_class", "ID"]
        selected = selected[keys].assign(displayed=True)
        df = df.merge(selected, on=keys, how="left", validate="one_to_one", sort=False)
        df["displayed"] = df["displayed"].eq(True)
        return df
    if name == "aucell_cells":
        metadata = read_table("results/gse162170/cell_programs/GSE4a_cell_coordinates_and_labels.tsv")
        scores = df.pivot(index="cell", columns="target_program", values="AUC").add_prefix("AUC_")
        return metadata.merge(scores, on="cell", how="left", validate="one_to_one", sort=False)
    if name == "erythroid_coordinates":
        trajectory = read_table("results/gse194122/trajectory/erythroid_coordinates_and_pseudotime.tsv.gz")
        return df.merge(trajectory[["cell_id", "GEX_pseudotime"]], on="cell_id",
                        how="left", validate="one_to_one", sort=False)
    raise ValueError(f"Unknown source transformation: {name}")


def load_source_sheet(entry):
    if "case" in entry:
        case = entry["case"]
        if case not in _case_views:
            with tempfile.TemporaryDirectory(prefix="beeline-case-") as directory:
                subprocess.run(r_function_command(
                    "code/functions/beeline_case_views.R", "write_beeline_case",
                    repo=PROJECT_ROOT, case=case, output=directory), env=task_environment(), check=True)
                _case_views[case] = {
                    name: pd.read_csv(os.path.join(directory, name + ".tsv"), sep="\t")
                    for name in ("truth", "edges", "summary")
                }
        df = _case_views[case][entry["view"]].copy()
    else:
        df = read_table(entry["path"])
    if entry.get("nonempty") and df.empty:
        raise ValueError(f"{entry['sheet']}: no data rows")
    if entry.get("identity"):
        column, expected = entry["identity"]
        if df.empty or df[column].isna().any() or set(df[column]) != {expected}:
            raise ValueError(f"{entry['sheet']}: expected {column}={expected} in {entry['path']}")
    if entry.get("transform"):
        df = transform_source(df, entry["transform"])
    if entry.get("columns"):
        df = df[entry["columns"]]
    for column, labels in entry.get("display_labels", {}).items():
        df[column] = df[column].replace(labels)
    if entry.get("nonempty") and df.empty:
        raise ValueError(f"{entry['sheet']}: no data rows")
    return _format_columns(df)


def write_workbook(xlsx_path, index, sheets):
    wb = Workbook(write_only=True)
    wb.properties.creator = "multiCSN"
    wb._fonts[0] = Font(name=FONT_NAME, size=11)
    used = set()
    header_fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL)
    header_font = Font(name=FONT_NAME, bold=True)
    for name, df in [("Table information", index)] + sheets:
        sheet_name = _sanitize_sheet_name(name)
        base = sheet_name
        suffix = 1
        while sheet_name in used:
            suffix += 1
            sheet_name = _sanitize_sheet_name(f"{base}_{suffix}")
        used.add(sheet_name)
        ws = wb.create_sheet(sheet_name)
        sample = [tuple(df.columns)] + list(df.head(100).itertuples(index=False, name=None))
        for col in range(len(df.columns)):
            width = max((len(str(row[col])) for row in sample if not pd.isna(row[col])), default=0)
            ws.column_dimensions[get_column_letter(col + 1)].width = min(max(width + 1, 10), 50)
        if sheet_name == "Table information":
            for column, width in {"A": 40, "B": 36, "C": 100}.items():
                ws.column_dimensions[column].width = width
        if sheet_name == "Fig3a_TF_paired_effects":
            ws.freeze_panes = "A2"
            ws.row_dimensions[1].height = 32
        header = []
        for value in df.columns:
            cell = WriteOnlyCell(ws, value=value)
            cell.fill = header_fill
            cell.font = header_font
            if sheet_name == "Fig3a_TF_paired_effects":
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            header.append(cell)
        ws.append(header)
        values = df.replace({float("inf"): "inf", -float("inf"): "-inf"}).astype(object)
        values = values.where(pd.notna(values), None)
        for row_number, row in enumerate(values.itertuples(index=False, name=None), 2):
            if sheet_name not in {"Table information", "Fig3a_TF_paired_effects"}:
                ws.append(row)
                continue
            ws.row_dimensions[row_number].height = 45 if sheet_name == "Table information" else 30
            cells = []
            for col, value in enumerate(row):
                cell = WriteOnlyCell(ws, value=value)
                if sheet_name == "Table information" and col == 2:
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                elif sheet_name == "Fig3a_TF_paired_effects":
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                cells.append(cell)
            ws.append(cells)
        log(f"{sheet_name}: {len(df):,} rows")
    destination = Path(xlsx_path)
    with tempfile.TemporaryDirectory(prefix=".source-data-", dir=destination.parent) as directory:
        candidate = Path(directory) / destination.name
        wb.save(candidate)
        os.replace(candidate, destination)


def main():
    _case_views.clear()
    os.makedirs(OUT_DIR, exist_ok=True)
    from export_sensitivity_tables import export_sensitivity_tables
    log(f"Building Source Data from {len(SOURCE_SHEETS)} result tables")

    index_rows = []
    sheets = []
    missing = []
    for entry in SOURCE_SHEETS:
        try:
            df = load_source_sheet(entry)
        except FileNotFoundError:
            missing.append(entry.get("path", entry.get("case")))
            log(f"{entry['sheet']}: source table not found", "fail")
            continue
        sheets.append((entry["sheet"], df))
        index_rows.append({
            "Figure": entry["figure"],
            "Sheet": _sanitize_sheet_name(entry["sheet"]),
            "Content": entry["content"],
            "Rows": len(df),
            "Columns": df.shape[1],
        })

    if missing:
        log(f"{len(missing)} source tables are missing; fix the list before shipping", "fail")
        for path in missing:
            log(f"  missing: {path}", "fail")
        return 1

    source_index = pd.DataFrame(index_rows).sort_values(
        "Figure", key=lambda s: s.map(_figure_sort_key), kind="stable"
    ).reset_index(drop=True)
    export_sensitivity_tables()
    write_workbook(
        SOURCE_XLSX,
        source_index,
        sheets,
    )
    log(f"Source Data written: {SOURCE_XLSX}", "info")

    return 0


def _figure_sort_key(label):
    """Sort figure labels as Figure 2 < Figure 3 < ... < Figure S1 < Figure S11."""
    text = str(label)
    main = re.match(r"Figure (\d+)", text)
    if main:
        return (0, int(main.group(1)), 0, text)
    supp = re.search(r"Figure S(\d+)", text)
    if supp:
        return (1, int(supp.group(1)), 0, text)
    return (2, 0, 0, text)


if __name__ == "__main__":
    sys.exit(main())
